"""
ICHITA PPTX Template Helpers
=============================
Reusable utilities for building presentations using ICHITA PowerPoint templates.
Follows ICHITA Brand Identity Guidelines V1.0.

Template: "New Identity" (16:9, 13.33" x 7.50")
Path: D:/Doccument/New Identity/Powerpoint Template.pptx

Layout Map:
  0: Title Slide       - idx 0=CENTER_TITLE, 1=SUBTITLE
  1: Title and Content  - idx 0=TITLE, 1=OBJECT(content)
  2: Section Header     - idx 0=TITLE, 1=BODY(text)
  3: Two Content        - idx 0=TITLE, 1=OBJECT(left), 2=OBJECT(right)
  4: Comparison          - idx 0=TITLE, 1=BODY(label-L), 2=OBJECT(content-L), 3=BODY(label-R), 4=OBJECT(content-R)
  5: Title Only          - idx 0=TITLE
  6: Blank               - (no content placeholders, only date/footer/slide#)
  7: Content with Caption - idx 0=TITLE, 1=OBJECT(content-R), 2=BODY(caption-L)
  8: Picture with Caption - idx 0=TITLE, 1=PICTURE, 2=BODY(caption-L)
  9: Title and Vertical Text
  10: Vertical Title and Text

Created: 2026-02-06
Updated: 2026-02-07 — Full brand compliance pass (12 fixes)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR_TYPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from lxml import etree
import copy
import os

# === ICHITA Brand Constants (from Brand Identity Guidelines V1.0) ===

TEMPLATE_PATH = "D:/Doccument/New Identity/Powerpoint Template.pptx"

# Brand Colors — exact values from guidelines
ICHITA_BLUE = RGBColor(0x29, 0x78, 0xFF)   # #2978FF — Primary accent
BLUE_LIGHT = RGBColor(0x82, 0xB0, 0xFF)    # #82B0FF — Secondary accent
BLUE_GREY_01 = RGBColor(0xCF, 0xD9, 0xDB)  # #CFD9DB — Light background
BLUE_GREY_02 = RGBColor(0x78, 0x8F, 0x9C)  # #788F9C — Muted text
BLUE_GREY_03 = RGBColor(0x26, 0x33, 0x38)  # #263338 — Primary text, dark bg
BLUE_BLACK = RGBColor(0x17, 0x1C, 0x21)    # #171C21 — Darkest background
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Legacy aliases for backward compatibility
BLUE = ICHITA_BLUE
DARK_TEAL = BLUE_GREY_03
NAVY = BLUE_BLACK
LIGHT_GRAY = BLUE_GREY_01
DARK_GRAY = BLUE_GREY_03

# Functional colors
GREEN = RGBColor(0x34, 0xA8, 0x53)      # #34A853 — Positive/success
RED = RGBColor(0xE8, 0x3E, 0x3E)        # #E83E3E — Warning/negative
ORANGE = RGBColor(0xFF, 0xA0, 0x00)     # #FFA000 — Emphasis (sparingly)
ALT_ROW = RGBColor(0xF0, 0xF4, 0xF5)    # Very light blue-grey for tables

# Fonts — ICHITA Brand Typography
FONT_HEADING = "Aeonik"                  # Headlines, titles
FONT_BODY = "Aeonik"                     # Body copy
FONT_DISPLAY = "Betatron"                # Large numbers, statistics
FONT_THAI = "TH Sarabun New"             # Thai text
FONT_FAMILY = "Aeonik"                   # Legacy alias

# Slide dimensions (16:9)
SLIDE_WIDTH = Inches(13.33)
SLIDE_HEIGHT = Inches(7.50)

# Content area constants (below title, above footer)
CONTENT_LEFT = Inches(0.92)
CONTENT_TOP = Inches(2.0)
CONTENT_WIDTH = Inches(11.5)
CONTENT_BOTTOM = Inches(6.5)

# Frame layout — logo banner occupies top-left ~3.3" x 0.55"
# Title for content slides: positioned to the RIGHT of the logo
TITLE_LEFT = Inches(3.8)       # Start after logo banner
TITLE_TOP = Inches(0.15)       # Near top of slide
TITLE_WIDTH = Inches(8.6)      # Fill remaining width
TITLE_HEIGHT = Inches(0.9)     # Enough room for long titles (was 0.6)

# Assets directory
ASSETS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
FRAME_PATH = os.path.join(ASSETS_DIR, "ichita-frame.png")
DARK_BG_PATH = os.path.join(ASSETS_DIR, "ichita-dark-bg.jpg")


# =====================================================================
# CORE FUNCTIONS
# =====================================================================

def load_template(path=TEMPLATE_PATH):
    """Load ICHITA template and return Presentation object."""
    prs = Presentation(path)
    # Remove sample slides that come with the template
    ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    while len(prs.slides) > 0:
        sldId = prs.slides._sldIdLst[0]
        rId = sldId.get(f'{ns}id')
        prs.part.drop_rel(rId)
        prs.slides._sldIdLst.remove(sldId)
    return prs


def apply_frame(slide):
    """Apply ICHITA content frame (PNG overlay) as the bottom-most shape.

    Frame has: logo banner top-left, white rounded content area, dark bottom strip.
    Call this FIRST after creating the slide, before adding content.
    """
    if os.path.exists(FRAME_PATH):
        pic = slide.shapes.add_picture(
            FRAME_PATH, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT
        )
        # Move to back (first in shape tree)
        sp = pic._element
        sp.getparent().remove(sp)
        slide.shapes._spTree.insert(2, sp)  # After background
    return slide


def apply_dark_bg(slide):
    """Apply ICHITA dark background (JPEG) as the bottom-most shape.

    Dark Blue Grey 03 with ICHITA logo centered at bottom.
    Used for title slides, section headers, closing slides.
    """
    if os.path.exists(DARK_BG_PATH):
        pic = slide.shapes.add_picture(
            DARK_BG_PATH, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT
        )
        sp = pic._element
        sp.getparent().remove(sp)
        slide.shapes._spTree.insert(2, sp)
    return slide


def reposition_title(slide, left=None, top=None, width=None, height=None):
    """Reposition the title placeholder to work with the frame layout.

    Default: mid-right position (right of logo banner).
    """
    left = left or TITLE_LEFT
    top = top or TITLE_TOP
    width = width or TITLE_WIDTH
    height = height or TITLE_HEIGHT

    try:
        title_ph = slide.placeholders[0]
        title_ph.left = left
        title_ph.top = top
        title_ph.width = width
        title_ph.height = height
    except (KeyError, IndexError):
        pass
    return slide


# =====================================================================
# STYLING HELPERS (NEW)
# =====================================================================

def _style_title_ph(slide, color=None, size=Pt(28), bold=True,
                    alignment=PP_ALIGN.LEFT):
    """Apply Aeonik font and brand styling to title placeholder."""
    color = color or BLUE_GREY_03
    try:
        title_ph = slide.placeholders[0]
        for p in title_ph.text_frame.paragraphs:
            p.alignment = alignment
            for run in p.runs:
                run.font.name = FONT_HEADING
                run.font.size = size
                run.font.bold = bold
                run.font.color.rgb = color
    except (KeyError, IndexError):
        pass


def add_accent_bar(slide, y=None, left=None, width=Inches(2.5)):
    """Add Ichita Blue accent bar (thin horizontal line) below title."""
    y = y or Inches(1.05)
    left = left or TITLE_LEFT
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, y, width, Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ICHITA_BLUE
    shape.line.fill.background()
    return shape


def add_footer(slide, text="www.ichita.co.th"):
    """Add footer text on the dark bottom strip of the frame."""
    txBox = slide.shapes.add_textbox(
        Inches(0.92), Inches(7.05), Inches(11.5), Inches(0.3)
    )
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    for run in p.runs:
        run.font.name = FONT_BODY
        run.font.size = Pt(9)
        run.font.color.rgb = WHITE
    return txBox


def _set_cell_border(cell, side, color=BLUE_GREY_01, width=Pt(0.5)):
    """Set border on one side of a table cell.

    Args:
        side: 'lnL', 'lnR', 'lnT', 'lnB'
        color: RGBColor object
        width: border width in EMU (use Pt() for convenience)
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    ln = tcPr.find(qn(f'a:{side}'))
    if ln is not None:
        tcPr.remove(ln)
    ln = etree.SubElement(tcPr, qn(f'a:{side}'))
    ln.set('w', str(int(width)))
    ln.set('cap', 'flat')
    ln.set('cmpd', 'sng')
    solidFill = etree.SubElement(ln, qn('a:solidFill'))
    srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
    srgbClr.set('val', str(color))


def _add_table_borders(table, color=BLUE_GREY_01, width=Pt(0.5)):
    """Add light borders to all cells of a table."""
    for row_idx in range(len(table.rows)):
        for col_idx in range(len(table.columns)):
            cell = table.cell(row_idx, col_idx)
            for side in ['lnL', 'lnR', 'lnT', 'lnB']:
                _set_cell_border(cell, side, color, width)


def _format_bullets(tf, bullets):
    """Format bullet list with visual hierarchy.

    Bullets starting with '  ' (two spaces) are treated as sub-bullets
    with smaller font and muted color.
    """
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        is_sub = bullet.startswith("  ")
        p.text = bullet.lstrip() if is_sub else bullet
        if is_sub:
            p.level = 1
        for run in p.runs:
            run.font.name = FONT_BODY
            if is_sub:
                run.font.size = Pt(16)
                run.font.color.rgb = BLUE_GREY_02
            else:
                run.font.size = Pt(20)
                run.font.color.rgb = BLUE_GREY_03


# =====================================================================
# LAYOUT FUNCTIONS
# =====================================================================

def add_title_slide(prs, title, subtitle=""):
    """Layout 0: Title Slide with dark background.

    White title text (Aeonik 44pt), Blue Light subtitle (Aeonik 20pt).
    Both centered on dark bg.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    apply_dark_bg(slide)

    # Title — white on dark bg
    title_ph = slide.placeholders[0]
    title_ph.text = title
    title_ph.left = Inches(1.5)
    title_ph.top = Inches(1.5)
    title_ph.width = Inches(10)
    title_ph.height = Inches(2.5)
    for p in title_ph.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.name = FONT_HEADING
            run.font.size = Pt(44)
            run.font.bold = True
            run.font.color.rgb = WHITE

    # Subtitle — Blue Light on dark bg
    if subtitle:
        sub_ph = slide.placeholders[1]
        sub_ph.text = subtitle
        sub_ph.left = Inches(1.5)
        sub_ph.top = Inches(4.0)
        sub_ph.width = Inches(10)
        sub_ph.height = Inches(2.5)
        for p in sub_ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = FONT_BODY
                run.font.size = Pt(20)
                run.font.color.rgb = BLUE_LIGHT

    return slide


def add_content_slide(prs, title, bullets):
    """Layout 1: Title and Content with frame background.

    Title positioned mid-right (next to logo), Aeonik 28pt.
    Bullets with visual hierarchy (L1: 20pt dark, L2: 16pt muted).
    Includes accent bar and footer.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_frame(slide)
    slide.shapes.title.text = title
    reposition_title(slide)
    _style_title_ph(slide)
    add_accent_bar(slide)

    # Reposition content area below title + accent bar
    content_ph = slide.placeholders[1]
    content_ph.left = Inches(0.92)
    content_ph.top = Inches(1.3)
    content_ph.width = Inches(11.5)
    content_ph.height = Inches(5.2)

    _format_bullets(content_ph.text_frame, bullets)

    add_footer(slide)
    return slide


def add_section_slide(prs, title, subtitle=""):
    """Layout 2: Section Header with dark background.

    White title text (Aeonik 40pt bold), Blue Light subtitle (Aeonik 18pt).
    Both centered.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    apply_dark_bg(slide)

    # Title — white on dark bg
    title_ph = slide.placeholders[0]
    title_ph.left = Inches(1.5)
    title_ph.top = Inches(1.5)
    title_ph.width = Inches(10)
    title_ph.height = Inches(3.0)
    title_ph.text = title
    for p in title_ph.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.name = FONT_HEADING
            run.font.size = Pt(40)
            run.font.bold = True
            run.font.color.rgb = WHITE

    # Subtitle — Blue Light
    if subtitle:
        sub_ph = slide.placeholders[1]
        sub_ph.left = Inches(1.5)
        sub_ph.top = Inches(4.5)
        sub_ph.width = Inches(10)
        sub_ph.height = Inches(2.0)
        sub_ph.text = subtitle
        for p in sub_ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.name = FONT_BODY
                run.font.size = Pt(18)
                run.font.color.rgb = BLUE_LIGHT

    return slide


def add_two_content_slide(prs, title, left_bullets, right_bullets):
    """Layout 3: Two Content with frame. Title mid-right.

    Includes accent bar and footer. Both columns formatted with hierarchy.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[3])
    apply_frame(slide)
    slide.shapes.title.text = title
    reposition_title(slide)
    _style_title_ph(slide)
    add_accent_bar(slide)

    # Left content
    _format_bullets(slide.placeholders[1].text_frame, left_bullets)

    # Right content
    _format_bullets(slide.placeholders[2].text_frame, right_bullets)

    add_footer(slide)
    return slide


def add_comparison_slide(prs, title, left_label, left_bullets,
                         right_label, right_bullets):
    """Layout 4: Comparison with frame. Title mid-right.

    Labels styled as category headers. Includes accent bar and footer.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[4])
    apply_frame(slide)
    slide.shapes.title.text = title
    reposition_title(slide)
    _style_title_ph(slide)
    add_accent_bar(slide)

    # Left label
    slide.placeholders[1].text = left_label
    for p in slide.placeholders[1].text_frame.paragraphs:
        for run in p.runs:
            run.font.name = FONT_HEADING
            run.font.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = ICHITA_BLUE

    # Left content
    _format_bullets(slide.placeholders[2].text_frame, left_bullets)

    # Right label
    slide.placeholders[3].text = right_label
    for p in slide.placeholders[3].text_frame.paragraphs:
        for run in p.runs:
            run.font.name = FONT_HEADING
            run.font.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = ICHITA_BLUE

    # Right content
    _format_bullets(slide.placeholders[4].text_frame, right_bullets)

    add_footer(slide)
    return slide


def add_blank_slide(prs, with_frame=True):
    """Layout 6: Blank slide. Optionally with frame."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if with_frame:
        apply_frame(slide)
        add_footer(slide)
    return slide


def add_title_only_slide(prs, title):
    """Layout 5: Title Only with frame. Title mid-right.

    Includes accent bar and footer. Large content area below for custom shapes.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    apply_frame(slide)
    slide.shapes.title.text = title
    reposition_title(slide)
    _style_title_ph(slide)
    add_accent_bar(slide)
    add_footer(slide)
    return slide


# =====================================================================
# TABLE HELPERS
# =====================================================================

def add_table(slide, rows, cols, data, left=None, top=None, width=None, height=None,
              header_color=BLUE, header_font_color=WHITE):
    """Add a native PowerPoint table to a slide.

    Args:
        data: list of lists, first row = headers
        left/top/width/height: position in Inches (defaults to centered content area)
    """
    if left is None:
        left = Inches(0.92)
    if top is None:
        top = Inches(2.0)
    if width is None:
        width = Inches(11.5)
    if height is None:
        height = Inches(0.4 * rows)

    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    for row_idx, row_data in enumerate(data):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(cell_text)

            # Format header row
            if row_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_color
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = header_font_color
                        run.font.bold = True
                        run.font.size = Pt(14)
                        run.font.name = FONT_FAMILY
            else:
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(12)
                        run.font.name = FONT_FAMILY

    # Add light borders
    _add_table_borders(table, BLUE_GREY_01, Pt(0.5))

    return table_shape


def add_styled_table(slide, data, left=None, top=None, width=None, height=None,
                     header_color=BLUE_GREY_03, header_text_color=WHITE,
                     alt_row_color=ALT_ROW, col_widths=None):
    """Add a professionally styled table with alternating rows and borders.

    Args:
        data: list of lists, first row = headers
        col_widths: list of Emu/Inches values for column widths
    """
    rows = len(data)
    cols = len(data[0]) if data else 0
    left = left or CONTENT_LEFT
    top = top or Inches(2.3)
    width = width or CONTENT_WIDTH
    height = height or Inches(0.45 * rows)

    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w

    for row_idx, row_data in enumerate(data):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(cell_text)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            if row_idx == 0:
                # Header row
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_color
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for run in p.runs:
                        run.font.color.rgb = header_text_color
                        run.font.bold = True
                        run.font.size = Pt(13)
                        run.font.name = FONT_HEADING
            else:
                # Data rows with alternating color
                if row_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = alt_row_color
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = WHITE
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(12)
                        run.font.name = FONT_BODY
                        run.font.color.rgb = BLUE_GREY_03

    # Add cell borders
    _add_table_borders(table, BLUE_GREY_01, Pt(0.5))

    return table_shape


def set_cell_fill(table, row, col, color):
    """Set background color for a specific cell."""
    cell = table.cell(row, col)
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def set_cell_text_color(table, row, col, color, bold=False):
    """Set text color for a specific cell."""
    cell = table.cell(row, col)
    for p in cell.text_frame.paragraphs:
        for run in p.runs:
            run.font.color.rgb = color
            if bold:
                run.font.bold = True


# =====================================================================
# TEXT HELPERS
# =====================================================================

def add_textbox(slide, text, left, top, width, height,
                font_size=Pt(18), bold=False, color=DARK_GRAY,
                alignment=PP_ALIGN.LEFT, font_name=FONT_FAMILY):
    """Add a text box with formatted text."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = alignment
    for run in p.runs:
        run.font.size = font_size
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name
    return txBox


def add_big_number(slide, number_text, label_text, left, top,
                   width=Inches(4), height=Inches(2)):
    """Add a large impact number (Betatron) with label (Aeonik) below."""
    # Number — Betatron display font, Ichita Blue
    num_box = slide.shapes.add_textbox(left, top, width, Inches(1.2))
    tf = num_box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = number_text
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.name = FONT_DISPLAY
        run.font.size = Pt(54)
        run.font.bold = True
        run.font.color.rgb = ICHITA_BLUE

    # Label — Aeonik
    add_textbox(slide, label_text, left, top + Inches(1.2), width, Inches(0.5),
                font_size=Pt(16), color=BLUE_GREY_02, alignment=PP_ALIGN.CENTER,
                font_name=FONT_BODY)


def add_process_arrow(slide, x1, y1, x2, y2, color=BLUE):
    """Add an arrow connector shape."""
    connector = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        x1, y1, x2 - x1, Inches(0.4)
    )
    connector.fill.solid()
    connector.fill.fore_color.rgb = color
    connector.line.fill.background()
    return connector


def add_process_box(slide, text, left, top, width=Inches(2.5), height=Inches(1.2),
                    fill_color=BLUE, font_color=WHITE, font_size=Pt(14)):
    """Add a rounded rectangle process box with text."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    for run in p.runs:
        run.font.size = font_size
        run.font.bold = True
        run.font.color.rgb = font_color
        run.font.name = FONT_FAMILY
    return shape


def format_placeholder_text(placeholder, items, font_size=Pt(18), bold_first_word=False):
    """Format a placeholder's text frame with multiple items."""
    tf = placeholder.text_frame
    tf.clear()
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        for run in p.runs:
            run.font.size = font_size
            run.font.name = FONT_FAMILY
    return tf


def probe_template(path=TEMPLATE_PATH):
    """Print all layouts and their placeholders for debugging."""
    prs = Presentation(path)
    print(f"Template: {path}")
    print(f"Size: {prs.slide_width/914400:.2f}\" x {prs.slide_height/914400:.2f}\"")
    print(f"Layouts: {len(prs.slide_layouts)}")
    print(f"Existing slides: {len(prs.slides)}")
    print()

    for i, layout in enumerate(prs.slide_layouts):
        print(f"Layout {i}: \"{layout.name}\"")
        for ph in layout.placeholders:
            idx = ph.placeholder_format.idx
            ptype = ph.placeholder_format.type
            if idx < 10:  # Skip date/footer/slide# (10,11,12)
                print(f"  [{idx}] {ptype} \"{ph.name}\" "
                      f"({ph.width/914400:.1f}\"x{ph.height/914400:.1f}\")")
        print()


# =====================================================================
# CHART HELPERS
# =====================================================================

def add_column_chart(slide, categories, series_dict, left=None, top=None,
                     width=None, height=None, colors=None, has_legend=True):
    """Add a clustered column chart with data labels and styled axes.

    Args:
        categories: list of category labels
        series_dict: {'Series Name': [val1, val2, ...], ...}
        colors: list of RGBColor for each series
    """
    left = left or CONTENT_LEFT
    top = top or Inches(2.5)
    width = width or Inches(8)
    height = height or Inches(4)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series_dict.items():
        chart_data.add_series(name, values)

    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        left, top, width, height, chart_data
    )
    chart = chart_frame.chart
    chart.has_legend = has_legend
    if has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT_BODY
        chart.legend.font.size = Pt(10)

    if colors:
        for i, color in enumerate(colors):
            if i < len(chart.series):
                chart.series[i].format.fill.solid()
                chart.series[i].format.fill.fore_color.rgb = color

    # Data labels
    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.show_value = True
    data_labels.show_category_name = False
    data_labels.show_series_name = False
    data_labels.font.name = FONT_BODY
    data_labels.font.size = Pt(9)
    data_labels.font.color.rgb = BLUE_GREY_03
    data_labels.number_format = '#,##0.0'

    # Style axes
    cat_axis = chart.category_axis
    cat_axis.has_minor_gridlines = False
    cat_axis.tick_labels.font.name = FONT_BODY
    cat_axis.tick_labels.font.size = Pt(10)
    cat_axis.tick_labels.font.color.rgb = BLUE_GREY_03

    val_axis = chart.value_axis
    val_axis.has_minor_gridlines = False
    val_axis.tick_labels.font.name = FONT_BODY
    val_axis.tick_labels.font.size = Pt(10)
    val_axis.tick_labels.font.color.rgb = BLUE_GREY_02

    return chart


def add_pie_chart(slide, categories, values, left=None, top=None,
                  width=None, height=None, colors=None, has_legend=True):
    """Add a pie chart with percentage data labels."""
    left = left or Inches(2)
    top = top or Inches(2.5)
    width = width or Inches(8)
    height = height or Inches(4)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series('Values', values)

    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        left, top, width, height, chart_data
    )
    chart = chart_frame.chart
    chart.has_legend = has_legend
    if has_legend:
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT_BODY
        chart.legend.font.size = Pt(10)

    if colors:
        series = chart.series[0]
        for i, color in enumerate(colors):
            series.points[i].format.fill.solid()
            series.points[i].format.fill.fore_color.rgb = color

    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.show_percentage = True
    data_labels.show_value = False
    data_labels.show_category_name = True
    data_labels.font.name = FONT_BODY
    data_labels.font.size = Pt(10)

    return chart


def add_bar_chart(slide, categories, series_dict, left=None, top=None,
                  width=None, height=None, colors=None, has_legend=True):
    """Add a horizontal bar chart with data labels and styled axes."""
    left = left or CONTENT_LEFT
    top = top or Inches(2.5)
    width = width or Inches(8)
    height = height or Inches(4)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series_dict.items():
        chart_data.add_series(name, values)

    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED,
        left, top, width, height, chart_data
    )
    chart = chart_frame.chart
    chart.has_legend = has_legend
    if has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT_BODY
        chart.legend.font.size = Pt(10)

    if colors:
        for i, color in enumerate(colors):
            if i < len(chart.series):
                chart.series[i].format.fill.solid()
                chart.series[i].format.fill.fore_color.rgb = color

    # Data labels
    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.show_value = True
    data_labels.show_category_name = False
    data_labels.show_series_name = False
    data_labels.font.name = FONT_BODY
    data_labels.font.size = Pt(9)
    data_labels.font.color.rgb = BLUE_GREY_03

    # Style axes
    cat_axis = chart.category_axis
    cat_axis.has_minor_gridlines = False
    cat_axis.tick_labels.font.name = FONT_BODY
    cat_axis.tick_labels.font.size = Pt(10)
    cat_axis.tick_labels.font.color.rgb = BLUE_GREY_03

    val_axis = chart.value_axis
    val_axis.has_minor_gridlines = False
    val_axis.tick_labels.font.name = FONT_BODY
    val_axis.tick_labels.font.size = Pt(10)
    val_axis.tick_labels.font.color.rgb = BLUE_GREY_02

    return chart


# =====================================================================
# DIAGRAM HELPERS
# =====================================================================

def add_process_flow_h(slide, steps, y_center, box_width=None, box_height=Inches(1.0),
                       fill_color=ICHITA_BLUE, text_color=WHITE, arrow_color=BLUE_GREY_02):
    """Add a horizontal process flow: [Step1] -> [Step2] -> [Step3]

    Args:
        steps: list of step label strings
        y_center: top position of the boxes
    Returns:
        list of shape objects
    """
    n = len(steps)
    gap = Inches(0.15)
    arrow_w = Inches(0.6)
    total_arrows = (n - 1) * (gap * 2 + arrow_w)
    usable = CONTENT_WIDTH - total_arrows
    bw = box_width or (usable / n)
    x = CONTENT_LEFT
    shapes = []

    for i, step in enumerate(steps):
        box = add_process_box(slide, step, x, y_center, bw, box_height,
                              fill_color=fill_color, font_color=text_color)
        shapes.append(box)

        if i < n - 1:
            ax = x + bw + gap
            ay = y_center + box_height / 2
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                ax, ay - Inches(0.15), arrow_w, Inches(0.3)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = arrow_color
            arrow.line.fill.background()

        x += bw + gap * 2 + arrow_w

    return shapes


def add_connector_arrow(slide, start_x, start_y, end_x, end_y,
                        color=ICHITA_BLUE, width=Pt(2)):
    """Add a straight connector with arrow head."""
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        start_x, start_y, end_x, end_y
    )
    connector.line.color.rgb = color
    connector.line.width = width
    # Arrow head
    ln = connector.line._ln
    tailEnd = ln.makeelement(qn('a:tailEnd'), {})
    tailEnd.set('type', 'triangle')
    tailEnd.set('w', 'med')
    tailEnd.set('len', 'med')
    ln.append(tailEnd)
    return connector


def add_circle_shape(slide, left, top, diameter, fill_color=ICHITA_BLUE,
                     text='', text_color=WHITE, font_size=Pt(16)):
    """Add a circle with optional centered text (for numbered steps, icons)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()

    if text:
        tf = shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = str(text)
        for run in p.runs:
            run.font.size = font_size
            run.font.bold = True
            run.font.color.rgb = text_color
            run.font.name = FONT_HEADING
        # Vertical center
        bodyPr = tf._txBody.find(qn('a:bodyPr'))
        if bodyPr is not None:
            bodyPr.set('anchor', 'ctr')
    return shape


# =====================================================================
# KPI / Big Number Helpers (Betatron for numbers)
# =====================================================================

def add_kpi_card(slide, left, top, width, height,
                 value, label, value_color=ICHITA_BLUE,
                 label_color=BLUE_GREY_02, border_color=BLUE_GREY_01):
    """Add a KPI card: big number (Betatron) + label (Aeonik)."""
    # Card border
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.background()
    card.line.color.rgb = border_color
    card.line.width = Pt(1)

    # Value — Betatron display font
    val_box = slide.shapes.add_textbox(left, top + Inches(0.15), width, Inches(0.8))
    tf = val_box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.name = FONT_DISPLAY
        run.font.size = Pt(48)
        run.font.bold = True
        run.font.color.rgb = value_color

    # Label — Aeonik
    add_textbox(slide, label, left, top + Inches(1.0), width, Inches(0.4),
                font_size=Pt(14), color=label_color,
                alignment=PP_ALIGN.CENTER, font_name=FONT_BODY)

    return card


# =====================================================================
# RICH TEXT HELPERS
# =====================================================================

def add_rich_text_box(slide, left, top, width, height):
    """Add an empty text box for rich text formatting. Returns text_frame."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf


def add_rich_run(paragraph, text, font_name=FONT_BODY, size=Pt(14),
                 color=BLUE_GREY_03, bold=False, italic=False):
    """Add a styled run to an existing paragraph (for mixed formatting)."""
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return run


# =====================================================================
# HORIZONTAL DIVIDER
# =====================================================================

def add_divider(slide, y, left=None, width=None, color=BLUE_GREY_01, thickness=Pt(1)):
    """Add a horizontal divider line."""
    left = left or CONTENT_LEFT
    width = width or CONTENT_WIDTH
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        left, y, left + width, y
    )
    connector.line.color.rgb = color
    connector.line.width = thickness
    return connector


if __name__ == "__main__":
    probe_template()
